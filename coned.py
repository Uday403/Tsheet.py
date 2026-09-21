from __future__ import annotations

import csv
import io
import re
import zipfile
from copy import copy
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


MASTER_TEMPLATE = Path(__file__).with_name("master_template.xlsm")
PRISMA_SHEET_CANDIDATES = (
    "Prisma Export - Paste as values",
    "Paste Prisma Export",
    "Prisma Export",
)
TRAFFIC_SHEET = "Traffic_Doc"
MULTI_SHEET_CANDIDATES = (
    "Multi-Ad or Creative Rotation",
    "Multi-Ad",
    "Multi-Creative Rotation",
)

SUPPORTED_CREATIVE_EXTENSIONS = {
    ".zip", ".gif", ".jpg", ".jpeg", ".png", ".webp",
    ".html", ".htm", ".mp4", ".mov", ".m4v",
}


def _clean(value) -> str:
    return "" if value is None else str(value).strip()


def _norm(value) -> str:
    return re.sub(r"[^a-z0-9]+", "", _clean(value).lower())


def _words(value) -> set[str]:
    text = _clean(value).lower()
    text = re.sub(r"\.(gif|jpg|jpeg|png|webp|html|htm|zip|mp4|mov|m4v)$", "", text)
    text = re.sub(r"\b(v|ver|version)[-_ ]?\d+\b", " ", text)
    text = re.sub(r"\b(160x600|300x250|300x50|300x600|320x50|728x90|970x250|970x90|1x1)\b", " ", text)
    text = re.sub(r"\b(6|10|15|20|30|60|90)\s*(s|sec|secs|second|seconds)\b", " ", text)
    stop = {
        "coned", "con", "ed", "creative", "banner", "display", "video",
        "audio", "final", "new", "ny", "oru", "cec", "2026", "2025",
        "jpg", "jpeg", "png", "gif", "html", "zip", "mp4", "mov",
    }
    return {w for w in re.findall(r"[a-z0-9]+", text) if len(w) >= 3 and w not in stop}


def _dimension(text: str) -> str:
    m = re.search(r"(?<!\d)(\d{2,4})\s*[xX]\s*(\d{2,4})(?!\d)", _clean(text))
    return f"{m.group(1)}x{m.group(2)}" if m else ""


def _duration(text: str) -> str:
    s = _clean(text)
    patterns = (
        r"(?<!\d)(6|10|15|20|30|60|90)\s*(?:s|sec|secs|second|seconds)(?![a-z0-9])",
        r"[:_ -](0?6|10|15|20|30|60|90)x(?:nan)?(?:_|-|$)",
    )
    for p in patterns:
        m = re.search(p, s, flags=re.I)
        if m:
            return str(int(m.group(1)))
    return ""


def _language(text: str) -> str:
    upper = _clean(text).upper()
    tokens = re.findall(r"[A-Z0-9]+", upper)
    if any(t in {"SP", "SPA", "SPANISH", "ES"} for t in tokens):
        return "SP"
    if any(t in {"EN", "ENG", "ENGLISH"} for t in tokens):
        return "EN"
    return ""


def _channel(text: str) -> str:
    u = _clean(text).upper()
    if "CTV" in u or "CONNECTED TV" in u or "CONNECTED_TV" in u:
        return "CTV"
    if "VIDEO" in u or "OLV" in u or "PREROLL" in u or "INSTREAM" in u:
        return "Video"
    if "AUDIO" in u or "DRAD" in u:
        return "Audio"
    if "NATIVE" in u:
        return "Native"
    if "DISPLAY" in u or "BANNER" in u:
        return "Display"
    return ""


def _concept(text: str) -> str:
    """
    Generic ConEd theme extraction.

    No campaign theme names are hard-coded here. Theme compatibility is
    determined dynamically from meaningful words shared by Placement Name
    and Creative File Name.
    """
    return ""


def _normalize_match_token(token: str) -> str:
    """
    Lightweight normalization for naming comparisons only.
    Handles normal singular/plural variations without hard-coding campaign themes:
      Innovations -> innovation
      Reports -> report
      Tools -> tool
      Upgrades -> upgrade
    """
    token = str(token or "").lower().strip()

    if len(token) > 5 and token.endswith("ies"):
        token = token[:-3] + "y"
    elif len(token) > 5 and token.endswith("ses"):
        token = token[:-2]
    elif len(token) > 4 and token.endswith("s") and not token.endswith(("ss", "us", "is")):
        token = token[:-1]

    return token


def _name_tokens(text: str) -> set[str]:
    """
    Meaningful generic tokens from placements, creatives or UTMs.
    No campaign/theme names such as Innovation, Customer Tools, Commercial,
    Residential, etc. are hard-coded.

    CamelCase is split:
      CleanEnergyInnovation -> clean, energy, innovation
      PartnerInGrowth       -> partner, in, growth
    """
    value = Path(_clean(text)).stem

    value = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", value)
    value = re.sub(r"(?<=[A-Z])(?=[A-Z][a-z])", " ", value)
    value = value.lower()

    tokens = re.findall(r"[a-z0-9]+", value)

    # Only trafficking/technical boilerplate is ignored.
    stop = {
        "programmatic", "display", "banner", "video", "audio", "native",
        "prospect", "behavioral", "retargeting", "geo", "data", "cpm",
        "device", "local", "web", "traffic", "party", "cross", "oath",
        "tps", "coned", "cec", "corp", "nan", "creative", "final",
        "version", "dark", "light", "english", "spanish",
        "png", "jpg", "jpeg", "gif", "webp", "html", "htm",
        "mp4", "mov", "m4v", "mp3", "wav", "m4a", "aac", "ogg",
        "http", "https", "www", "com", "utm", "source", "medium",
        "campaign", "content", "term", "assembly", "cpc", "sec",
        "none", "broad",
    }

    result = set()
    for raw in tokens:
        token = _normalize_match_token(raw)

        if token in stop:
            continue
        if len(token) < 3:
            continue
        if token.isdigit():
            continue
        if re.fullmatch(r"\d{2,4}", token):
            continue
        if re.fullmatch(r"fy\d+", token):
            continue
        if re.fullmatch(r"\d+x\d+", token):
            continue

        result.add(token)

    return result


def _semantic_shared_tokens(left: str, right: str) -> set[str]:
    """
    Generic semantic-name overlap.

    Exact normalized tokens match first. We also allow a long token to contain
    another long token, which helps naming variations while remaining generic.
    """
    left_tokens = _name_tokens(left)
    right_tokens = _name_tokens(right)

    shared = left_tokens & right_tokens

    for a in left_tokens:
        for b in right_tokens:
            if a == b:
                continue
            if min(len(a), len(b)) < 6:
                continue
            if a in b or b in a:
                shared.add(a if len(a) <= len(b) else b)

    return shared


def _shared_name_score(placement_name: str, creative_name: str) -> tuple[int, set[str]]:
    """
    Score only meaningful naming relationships.
    Dimensions/language/channel are handled separately as hard filters and
    therefore cannot accidentally make the wrong creative family win.
    """
    shared = _semantic_shared_tokens(placement_name, creative_name)
    score = sum(max(3, min(len(token), 16)) for token in shared)
    return score, shared

def _season(text: str) -> str:
    low = _clean(text).lower()
    if "summer" in low:
        return "Summer"
    if "winter" in low:
        return "Winter"
    return ""


def _version(text: str) -> str:
    m = re.search(r"(?:^|[-_ ])(?:v|ver|version)[-_ ]?(\d+)(?:$|[-_ .])", _clean(text), re.I)
    return f"V{m.group(1)}" if m else ""


def _creative_names(files: Iterable) -> list[str]:
    names = []
    for uploaded in files or []:
        name = Path(uploaded.name).name
        if name.lower().endswith(".zip"):
            uploaded.seek(0)
            with zipfile.ZipFile(uploaded, "r") as zf:
                for member in zf.namelist():
                    if member.endswith("/"):
                        continue
                    base = Path(member).name
                    if not base or base.startswith(".") or base.startswith("__MACOSX"):
                        continue
                    if Path(base).suffix.lower() in SUPPORTED_CREATIVE_EXTENSIONS:
                        names.append(base)
        elif Path(name).suffix.lower() in SUPPORTED_CREATIVE_EXTENSIONS:
            names.append(name)
    return list(dict.fromkeys(names))


def _read_prisma(prisma_file):
    prisma_file.seek(0)
    raw = prisma_file.read()
    if isinstance(raw, str):
        text = raw
    else:
        text = raw.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))

    header_idx = None
    headers = None
    for i, row in enumerate(rows[:50]):
        normalized = [_norm(v) for v in row]
        if "placementname" in normalized:
            header_idx = i
            headers = [_clean(v) for v in row]
            break
    if header_idx is None:
        raise ValueError("Could not find the Placement Name header in the Prisma export.")

    records = []
    for excel_row, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        padded = row + [""] * max(0, len(headers) - len(row))
        rec = dict(zip(headers, padded[:len(headers)]))
        rec["_source_excel_row"] = excel_row
        placement = _get(rec, "Placement Name")
        row_type = _get(rec, "Row Type", "Type", "Package/Placement")
        if not placement:
            continue
        if _clean(row_type).lower() == "package" or placement.lower().startswith("package:"):
            continue
        records.append(rec)
    return rows, records


def _get(record, *aliases):
    wanted = {_norm(x) for x in aliases}
    for k, v in record.items():
        if _norm(k) in wanted:
            return _clean(v)
    return ""


def _placement_name(record) -> str:
    return _get(record, "Placement Name")


def _placement_attributes(record) -> dict:
    name = _placement_name(record)
    return {
        "placement": name,
        "dimension": _dimension(name) or _get(record, "Dimensions", "Dimension", "Creative Size"),
        "duration": _duration(name),
        "language": _language(name),
        "channel": _channel(name),
        "concept": _concept(name),
        "season": _season(name),
    }


def _creative_attributes(name: str) -> dict:
    return {
        "name": name,
        "dimension": _dimension(name),
        "duration": _duration(name),
        "language": _language(name),
        "channel": _channel(name),
        "concept": _concept(name),
        "season": _season(name),
        "version": _version(name),
        "words": _words(name),
    }


def _compatible_channel(placement_channel: str, creative_name: str) -> bool:
    ext = Path(creative_name).suffix.lower()
    if placement_channel in {"Video", "CTV"}:
        return ext in {".mp4", ".mov", ".m4v"}
    if placement_channel == "Audio":
        return ext in {".mp3", ".wav", ".m4a"}
    if placement_channel in {"Display", "Native", ""}:
        return ext not in {".mp4", ".mov", ".m4v", ".mp3", ".wav", ".m4a"}
    return True


def match_creatives(record, creative_names: list[str]) -> tuple[list[str], str]:
    """
    Generic ConEd placement -> creative matching.

    Technical attributes are HARD filters.
    Naming/family relationship is then selected dynamically.

    Important:
      - No campaign/theme list is maintained.
      - "Innovation" can match "CleanEnergyInnovation".
      - "Innovations" also matches "Innovation".
      - A future unseen naming family works the same way.
      - Multiple strongest creative variants become Multi-Ad.
      - Multiple technically valid creatives with no naming relationship are
        Ambiguous instead of being guessed.
    """
    p = _placement_attributes(record)
    placement_name = p["placement"]
    candidates = []

    for name in creative_names:
        c = _creative_attributes(name)

        if not _compatible_channel(p["channel"], name):
            continue

        if p["dimension"]:
            if not c["dimension"] or p["dimension"].lower() != c["dimension"].lower():
                continue

        if p["duration"]:
            if not c["duration"] or p["duration"] != c["duration"]:
                continue

        if p["language"]:
            if not c["language"] or p["language"] != c["language"]:
                continue

        if p["season"] and c["season"] != p["season"]:
            continue

        score, shared = _shared_name_score(placement_name, name)

        candidates.append({
            "name": name,
            "score": score,
            "shared": shared,
        })

    if not candidates:
        return [], "Unmatched"

    if len(candidates) == 1:
        return [candidates[0]["name"]], "Direct"

    best_score = max(item["score"] for item in candidates)

    # Never let same size/language alone decide among different creative sets.
    if best_score <= 0:
        return [], "Ambiguous"

    matches = sorted(
        {item["name"] for item in candidates if item["score"] == best_score},
        key=str.lower,
    )

    return matches, "Multi" if len(matches) > 1 else "Direct"

def _parse_urls(text: str) -> list[str]:
    urls = []
    for line in _clean(text).splitlines():
        line = line.strip()
        if not line:
            continue
        m = re.search(r"https?://\S+", line)
        urls.append(m.group(0).rstrip(",;") if m else line)
    return list(dict.fromkeys(urls))


def _url_attributes(url: str) -> dict:
    from urllib.parse import unquote_plus
    decoded = unquote_plus(url.replace("&amp;", "&"))
    return {
        "url": url,
        "dimension": _dimension(decoded),
        "duration": _duration(decoded),
        "language": _language(decoded) or ("SP" if "/es/" in decoded.lower() else ""),
        "channel": _channel(decoded),
        "concept": _concept(decoded),
        "season": _season(decoded),
        "words": _words(decoded),
    }


def match_url(record, urls: list[str], creative_name: str = "") -> tuple[str, str]:
    """
    Generic URL/UTM matching.

    Uses technical signals plus semantic naming from the placement/creative
    against the decoded URL/UTM. Singular/plural variations are normalized,
    so placement "Innovations" can map to utm_term=ce_innovation_....

    If two URLs remain equally valid, return Ambiguous instead of guessing.
    """
    from urllib.parse import unquote_plus

    p = _placement_attributes(record)
    c = _creative_attributes(creative_name) if creative_name else {}
    placement_text = p["placement"]
    creative_text = creative_name or ""

    ranked = []

    for url in urls:
        u = _url_attributes(url)
        decoded_url = unquote_plus(url.replace("&amp;", "&"))

        # Explicit technical conflicts eliminate the URL.
        conflict = False
        for key in ("language", "channel", "dimension", "duration", "season"):
            expected = p.get(key) or c.get(key, "")
            actual = u.get(key, "")
            if expected and actual and expected != actual:
                conflict = True
                break

        if conflict:
            continue

        score = 0

        # Strong technical evidence.
        for key, weight in (
            ("language", 60),
            ("channel", 45),
            ("dimension", 40),
            ("duration", 40),
            ("season", 35),
        ):
            expected = p.get(key) or c.get(key, "")
            actual = u.get(key, "")
            if expected and actual and expected == actual:
                score += weight

        # Dynamic family/theme evidence from both placement and creative.
        placement_shared = _semantic_shared_tokens(placement_text, decoded_url)
        creative_shared = _semantic_shared_tokens(creative_text, decoded_url) if creative_text else set()

        score += sum(max(5, min(len(t), 18)) for t in placement_shared)
        score += sum(max(7, min(len(t), 20)) for t in creative_shared)

        ranked.append({
            "score": score,
            "url": url,
            "placement_shared": placement_shared,
            "creative_shared": creative_shared,
        })

    if not ranked:
        return "", "Unmatched"

    best_score = max(item["score"] for item in ranked)
    best = [item for item in ranked if item["score"] == best_score]

    if len(best) > 1:
        return "", "Ambiguous"

    return best[0]["url"], "Matched"

def build_ad_name(record, matches: list[str]) -> str:
    """
    ConEd rule: Placement Name = Ad Name exactly.
    """
    return _placement_name(record)


def preview_coned_setup(prisma_file, creative_files, urls_text: str = "") -> dict:
    _, records = _read_prisma(prisma_file)
    creatives = _creative_names(creative_files)
    urls = _parse_urls(urls_text)
    rows = []
    warnings = []

    for rec in records:
        matches, destination = match_creatives(rec, creatives)
        ad_name = build_ad_name(rec, matches)

        url_status = "Not supplied"
        matched_url = ""
        if urls:
            # For Multi, URL can be creative-specific. Preview the first match;
            # generation resolves every creative independently.
            matched_url, url_status = match_url(rec, urls, matches[0] if matches else "")

        attrs = _placement_attributes(rec)
        row = {
            "Placement Name": attrs["placement"],
            "Ad Name": ad_name,
            "Channel": attrs["channel"],
            "Concept": attrs["concept"],
            "Season": attrs["season"],
            "Size/Duration": attrs["dimension"] or (attrs["duration"] + "s" if attrs["duration"] else ""),
            "Destination": destination,
            "Creative Count": len(matches),
            "Creatives": " | ".join(matches),
            "URL Status": url_status,
            "URL": matched_url,
        }
        rows.append(row)

        if destination in {"Unmatched", "Ambiguous"}:
            warnings.append(f"{destination} creative: {attrs['placement']}")
        if urls and url_status != "Matched":
            warnings.append(f"{url_status} URL: {attrs['placement']}")

    return {
        "rows": rows,
        "warnings": warnings,
        "placement_count": len(records),
        "creative_count": len(creatives),
    }


def _find_sheet(workbook, candidates):
    normalized = {_norm(s): s for s in workbook.sheetnames}
    for name in candidates:
        if _norm(name) in normalized:
            return workbook[normalized[_norm(name)]]
    return None


def _find_header_row(sheet, required_phrases, max_rows=30):
    for r in range(1, min(sheet.max_row, max_rows) + 1):
        vals = [_clean(sheet.cell(r, c).value) for c in range(1, sheet.max_column + 1)]
        normed = [_norm(v) for v in vals]
        if all(any(_norm(p) in h for h in normed) for p in required_phrases):
            return r
    return None


def _header_map(sheet, row):
    return {_norm(sheet.cell(row, c).value): c for c in range(1, sheet.max_column + 1) if _clean(sheet.cell(row, c).value)}


def _col(hmap, *names, required=False):
    for n in names:
        nn = _norm(n)
        if nn in hmap:
            return hmap[nn]
    for h, c in hmap.items():
        if any(_norm(n) in h for n in names):
            return c
    if required:
        raise ValueError("Missing template column: " + " / ".join(names))
    return None


def _copy_row_style(sheet, source_row, target_row, max_col):
    for c in range(1, max_col + 1):
        src = sheet.cell(source_row, c)
        dst = sheet.cell(target_row, c)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def _paste_prisma(workbook, raw_rows):
    sheet = _find_sheet(workbook, PRISMA_SHEET_CANDIDATES)
    if sheet is None:
        return
    for row in sheet.iter_rows():
        for cell in row:
            cell.value = None
    for r, values in enumerate(raw_rows, 1):
        for c, value in enumerate(values, 1):
            sheet.cell(r, c).value = value


def _date_value(record, start=True):
    if start:
        return _get(record, "Flight start date", "Start Date", "Placement Start Date")
    return _get(record, "Flight end date", "End Date", "Placement End Date")


def _site_name(record):
    return _get(record, "Site Name", "Media Outlet", "Supplier Name", "Vendor") or "Oath Programmatic"


def _placement_id(record):
    return _get(record, "Ad server ID", "Ad Server ID", "DCM Placement ID", "Placement ID")


def generate_coned_tsheet(
    prisma_file,
    creative_files,
    urls_text: str = "",
    override_start_date=None,
    override_end_date=None,
    template_path=None,
) -> tuple[bytes, list[str]]:
    template = Path(template_path) if template_path else MASTER_TEMPLATE
    if not template.exists():
        raise FileNotFoundError(f"ConEd template not found: {template}")

    raw_rows, records = _read_prisma(prisma_file)
    creative_names = _creative_names(creative_files)
    urls = _parse_urls(urls_text)

    wb = load_workbook(template, keep_vba=True)
    _paste_prisma(wb, raw_rows)

    traffic = wb[TRAFFIC_SHEET]
    tr_header = _find_header_row(traffic, ("Placement Name", "Ad Name", "Creative File Name"))
    if tr_header is None:
        raise ValueError("Could not locate Traffic_Doc headers.")
    th = _header_map(traffic, tr_header)
    first_tr = tr_header + 1

    cols = {
        "additional": _col(th, "Additional Pixels"),
        "site": _col(th, "Site Name"),
        "id": _col(th, "DCM Placement ID", "Ad server ID"),
        "placement": _col(th, "Placement Name", required=True),
        "dimension": _col(th, "Dimensions", "Dimension"),
        "video": _col(th, "Video Length"),
        "vast": _col(th, "Vast/Vpaid", "VAST/VPAID"),
        "ad": _col(th, "Ad Name", required=True),
        "notes": _col(th, "Trafficking Notes"),
        "action": _col(th, "Action"),
        "creative": _col(th, "Creative File Name", required=True),
        "studio": _col(th, "Studio Creative"),
        "rotation": _col(th, "Rotation"),
        "start": _col(th, "Start Date"),
        "end": _col(th, "End Date"),
        "url": _col(th, "Click through URL", "Clickthrough URL"),
    }

    # Preserve one template style row and clear old data values.
    max_col = max(traffic.max_column, max(c for c in cols.values() if c))
    for r in range(first_tr, max(traffic.max_row, first_tr) + 1):
        for c in range(1, max_col + 1):
            traffic.cell(r, c).value = None

    multi = _find_sheet(wb, MULTI_SHEET_CANDIDATES)
    mh = {}
    first_multi = None
    if multi is not None:
        mr = _find_header_row(multi, ("Creative File Name", "Rotation", "Start Date", "End Date"))
        if mr:
            mh = _header_map(multi, mr)
            first_multi = mr + 1
            for r in range(first_multi, max(multi.max_row, first_multi) + 1):
                for c in range(1, multi.max_column + 1):
                    multi.cell(r, c).value = None

    warnings = []
    multi_row = first_multi
    written_multi = set()

    for idx, rec in enumerate(records):
        row = first_tr + idx
        if row != first_tr:
            _copy_row_style(traffic, first_tr, row, max_col)

        placement = _placement_name(rec)
        attrs = _placement_attributes(rec)
        matches, destination = match_creatives(rec, creative_names)
        ad_name = build_ad_name(rec, matches)
        start_date = override_start_date or _date_value(rec, True)
        end_date = override_end_date or _date_value(rec, False)

        # Resolve URL for direct creative. For Multi, URL is intentionally
        # creative-specific and written in Multi-Ad.
        direct_url = ""
        if len(matches) == 1 and urls:
            direct_url, url_status = match_url(rec, urls, matches[0])
            if url_status != "Matched":
                warnings.append(f"{url_status} URL: {placement}")

        values = {
            cols["site"]: _site_name(rec),
            cols["id"]: _placement_id(rec),
            cols["placement"]: placement,
            cols["dimension"]: attrs["dimension"] or ("1x1" if attrs["channel"] in {"Video", "CTV", "Audio"} else ""),
            cols["video"]: attrs["duration"],
            cols["ad"]: ad_name,
            cols["action"]: "New",
            cols["creative"]: matches[0] if len(matches) == 1 else "",
            cols["studio"]: "N" if len(matches) == 1 else "",
            cols["rotation"]: "100%" if len(matches) == 1 else "",
            cols["start"]: start_date,
            cols["end"]: end_date,
            cols["url"]: direct_url,
        }
        for c, v in values.items():
            if c:
                traffic.cell(row, c).value = v

        if destination in {"Unmatched", "Ambiguous"}:
            warnings.append(f"{destination} creative: {placement}")

        if len(matches) > 1:
            if multi is None or first_multi is None:
                warnings.append(f"Multi-Ad sheet not found for: {placement}")
                continue

            # Avoid duplicating the same ad block for repeated tactic placements.
            key = (ad_name, tuple(matches), start_date, end_date)
            if key in written_multi:
                continue
            written_multi.add(key)

            mc = {
                "ad": _col(mh, "Ad Name", "AD Name"),
                "action": _col(mh, "Action"),
                "creative": _col(mh, "Creative File Name", required=True),
                "studio": _col(mh, "Studio Creative"),
                "rotation": _col(mh, "Rotation", required=True),
                "start": _col(mh, "Start Date", required=True),
                "end": _col(mh, "End Date", required=True),
                "url": _col(mh, "Click through URL", "Clickthrough URL", required=True),
            }

            for creative in matches:
                if multi_row != first_multi:
                    _copy_row_style(multi, first_multi, multi_row, multi.max_column)

                creative_url = ""
                if urls:
                    creative_url, status = match_url(rec, urls, creative)
                    if status != "Matched":
                        warnings.append(f"{status} URL for {creative}: {placement}")

                vals = {
                    mc["ad"]: ad_name,
                    mc["action"]: "New",
                    mc["creative"]: creative,
                    mc["studio"]: "N",
                    mc["rotation"]: "Even",
                    mc["start"]: start_date,
                    mc["end"]: end_date,
                    mc["url"]: creative_url,
                }
                for c, v in vals.items():
                    if c:
                        multi.cell(multi_row, c).value = v
                multi_row += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue(), list(dict.fromkeys(warnings))
