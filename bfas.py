from __future__ import annotations

import io
import re
import zipfile
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Iterable
from urllib.parse import urlsplit


BASE_DIR = Path(__file__).resolve().parent
MASTER_TEMPLATE = BASE_DIR / "master_template.xlsm"

TRAFFIC_SHEET = "Traffic_Doc"
MULTI_SHEET = "Multi-Ad or Creative Rotation"

SUPPORTED_CREATIVE_EXTENSIONS = {
    ".jpg", ".jpeg", ".png", ".gif", ".webp",
    ".html", ".htm", ".mp4",
}

GENERIC_WORDS = {
    "bfas", "bfs", "creative", "creatives", "display", "olv",
    "final", "finals", "platform", "refresh", "ref", "fy",
    "q1", "q2", "q3", "q4", "new", "static", "banner", "banners",
    "jpg", "jpeg", "png", "gif", "webp", "html", "htm", "mp4",
}


def _clean(value) -> str:
    return "" if value is None else str(value).strip()


def _normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", _clean(value).lower())


def _words(value: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", _clean(value).lower())


def _meaningful_tokens(value: str) -> set[str]:
    tokens = set()
    for token in _words(value):
        if token in GENERIC_WORDS:
            continue
        if re.fullmatch(r"\d+", token):
            continue
        if re.fullmatch(r"\d{2,4}x\d{2,4}", token):
            continue
        if re.fullmatch(r"\d+s", token):
            continue
        if len(token) < 2:
            continue
        tokens.add(token)
    return tokens


def _extract_dimension(value: str) -> str:
    match = re.search(
        r"(?<!\d)(\d{2,4})\s*[xX×*]\s*(\d{2,4})(?!\d)",
        _clean(value),
    )
    return f"{match.group(1)}x{match.group(2)}" if match else ""


def _extract_video_length(value: str) -> str:
    match = re.search(r"(?<!\d)(6|10|15|20|30|60|90)\s*s(?:ec(?:ond)?s?)?(?!\d)", _clean(value), re.I)
    if match:
        return f"{match.group(1)}s"

    # Also catch names such as OLV-15s.
    match = re.search(r"(?<!\d)(6|10|15|20|30|60|90)s(?!\d)", _clean(value), re.I)
    return f"{match.group(1)}s" if match else ""


def _placement_size(value: str) -> str:
    return _extract_dimension(value) or _extract_video_length(value)


def _is_url(value: str) -> bool:
    try:
        parts = urlsplit(_clean(value))
        return parts.scheme in {"http", "https"} and bool(parts.netloc)
    except Exception:
        return False


def _creative_names_from_uploads(creative_files: Iterable) -> list[str]:
    """
    Read actual creative assets from direct uploads or ZIP files.
    Nested ZIPs/system files are ignored so they cannot create false ties.
    """
    names: list[str] = []

    for uploaded_file in creative_files or []:
        file_name = Path(uploaded_file.name).name
        suffix = Path(file_name).suffix.lower()

        if suffix == ".zip":
            uploaded_file.seek(0)
            with zipfile.ZipFile(uploaded_file, "r") as archive:
                for member in archive.namelist():
                    if member.endswith("/"):
                        continue

                    name = Path(member).name
                    if not name or name.startswith("."):
                        continue

                    if Path(name).suffix.lower() in SUPPORTED_CREATIVE_EXTENSIONS:
                        names.append(name)

        elif suffix in SUPPORTED_CREATIVE_EXTENSIONS:
            names.append(file_name)

    return list(dict.fromkeys(names))


def parse_placement_names(placement_text: str) -> list[str]:
    placements = []

    for raw_line in _clean(placement_text).splitlines():
        line = raw_line.strip()
        if not line:
            continue

        # Allows users to paste one Excel column.
        first_cell = line.split("\t")[0].strip()
        if first_cell and first_cell not in placements:
            placements.append(first_cell)

    return placements


def _creative_score(creative_name: str, placement_name: str) -> float:
    required_size = _placement_size(placement_name)
    creative_size = _placement_size(creative_name)

    # Dimension/video length is mandatory whenever placement supplies one.
    if required_size:
        if not creative_size:
            return -1000
        if creative_size.lower() != required_size.lower():
            return -1000

    score = 40.0 if required_size else 0.0

    placement_tokens = _meaningful_tokens(placement_name)
    creative_tokens = _meaningful_tokens(Path(creative_name).stem)
    overlap = placement_tokens & creative_tokens

    score += len(overlap) * 12

    # Stronger weight for useful version/set markers such as R4.
    for token in overlap:
        if re.fullmatch(r"[a-z]+\d+|\d+[a-z]+", token):
            score += 15

    # Long shared chunks help with phrases like "Engagement Creative".
    pnorm = _normalize(placement_name)
    for segment in re.split(r"[_\-\s]+", Path(creative_name).stem):
        snorm = _normalize(segment)
        if len(snorm) >= 6 and snorm in pnorm:
            score += 8

    return score


def match_creative(
    creative_names: list[str],
    placement_name: str,
) -> tuple[str, list[str]]:
    ranked = sorted(
        (
            (_creative_score(name, placement_name), name)
            for name in creative_names
        ),
        key=lambda item: (-item[0], item[1].lower()),
    )

    if not ranked or ranked[0][0] <= 0:
        return "", []

    best_score = ranked[0][0]
    tied = [name for score, name in ranked if score == best_score]

    if len(tied) > 1:
        return "", tied

    return ranked[0][1], []


def creative_set_key(creative_name: str) -> str:
    """
    Produce a readable creative-set key by removing size/length and generic
    filename noise. Example:
      BFS-2607-Engagement Creative-R4-300x250.jpg
    becomes approximately:
      2607 Engagement R4
    """
    stem = Path(creative_name).stem
    stem = re.sub(r"(?<!\d)\d{2,4}\s*[xX×*]\s*\d{2,4}(?!\d)", " ", stem)
    stem = re.sub(r"(?<!\d)(6|10|15|20|30|60|90)\s*s(?!\d)", " ", stem, flags=re.I)

    kept = []
    for token in re.findall(r"[A-Za-z0-9]+", stem):
        low = token.lower()
        if low in GENERIC_WORDS:
            continue
        kept.append(token)

    return " ".join(kept).strip() or stem.strip()


def parse_url_mapping(url_mapping_text: str) -> list[dict]:
    """
    Supported input:
      1) One URL only:
         https://...
         -> applies to every matched creative.

      2) Set keyword + URL:
         Engagement R4<TAB>https://...
         Adoption<TAB>https://...

      A pipe can also be used instead of a tab.
    """
    rows = []

    for raw_line in _clean(url_mapping_text).splitlines():
        line = raw_line.strip()
        if not line:
            continue

        if _is_url(line):
            rows.append({"key": "", "url": line})
            continue

        parts = re.split(r"\t+|\s*\|\s*", line, maxsplit=1)
        if len(parts) == 2 and _is_url(parts[1]):
            rows.append({
                "key": parts[0].strip(),
                "url": parts[1].strip(),
            })

    return rows


def match_url_for_creative(
    creative_name: str,
    url_rows: list[dict],
) -> tuple[str, str]:
    if not url_rows:
        return "", "missing"

    # A single bare URL applies to all creatives.
    if len(url_rows) == 1 and not url_rows[0]["key"]:
        return url_rows[0]["url"], "single"

    set_key = creative_set_key(creative_name)
    creative_text = f"{Path(creative_name).stem} {set_key}"
    creative_tokens = _meaningful_tokens(creative_text)

    scored = []
    for row in url_rows:
        key = row["key"]
        if not key:
            continue

        key_tokens = _meaningful_tokens(key)
        overlap = len(creative_tokens & key_tokens)

        score = overlap * 10
        if _normalize(key) and _normalize(key) in _normalize(creative_text):
            score += 30

        if score > 0:
            scored.append((score, row["url"], key))

    if not scored:
        return "", "missing"

    scored.sort(reverse=True)
    best_score = scored[0][0]
    best = [row for row in scored if row[0] == best_score]

    unique_urls = list(dict.fromkeys(row[1] for row in best))
    if len(unique_urls) != 1:
        return "", "ambiguous"

    return unique_urls[0], "mapped"


def preview_bfas_setup(
    placement_text: str,
    creative_files,
    url_mapping_text: str = "",
) -> dict:
    placements = parse_placement_names(placement_text)
    creative_names = _creative_names_from_uploads(creative_files)
    url_rows = parse_url_mapping(url_mapping_text)

    rows = []
    warnings = []

    for placement in placements:
        creative, tied = match_creative(creative_names, placement)
        size = _placement_size(placement)

        url = ""
        url_status = "missing"

        if creative:
            url, url_status = match_url_for_creative(creative, url_rows)

        if tied:
            warnings.append(
                f"Ambiguous creative match for '{placement}': "
                + ", ".join(tied)
            )
        elif not creative:
            warnings.append(
                f"No creative matched placement '{placement}'."
            )

        if creative and not url:
            if url_status == "ambiguous":
                warnings.append(
                    f"Ambiguous URL mapping for creative '{creative}'."
                )
            else:
                warnings.append(
                    f"No URL mapping found for creative '{creative}' "
                    f"(set: {creative_set_key(creative)})."
                )

        rows.append({
            "placement_name": placement,
            "ad_name": placement,
            "size": size,
            "creative": creative,
            "creative_set": creative_set_key(creative) if creative else "",
            "url": url,
            "url_status": url_status,
            "ambiguous_creatives": tied,
        })

    return {
        "rows": rows,
        "placements": rows,
        "creative_names": creative_names,
        "creative_count": len(creative_names),
        "matched_count": sum(1 for row in rows if row["creative"]),
        "unmatched_count": sum(1 for row in rows if not row["creative"]),
        "url_matched_count": sum(1 for row in rows if row["url"]),
        "url_unmatched_count": sum(1 for row in rows if not row["url"]),
        "warnings": warnings,
    }


def _find_header_row(sheet) -> int:
    for row in range(1, min(sheet.max_row, 40) + 1):
        values = {
            _normalize(sheet.cell(row=row, column=col).value)
            for col in range(1, min(sheet.max_column, 40) + 1)
        }
        if "placementname" in values and "adname" in values:
            return row
    raise ValueError("Unable to locate the Traffic_Doc header row.")


def _header_map(sheet, header_row: int) -> dict[str, int]:
    mapping = {}

    for col in range(1, sheet.max_column + 1):
        value = _normalize(sheet.cell(header_row, col).value)

        if value == "additionalpixels":
            mapping["additional_pixels"] = col
        elif value == "sitename":
            mapping["site"] = col
        elif value in {"dcmplacementid", "placementid"}:
            mapping["placement_id"] = col
        elif value == "placementname":
            mapping["placement"] = col
        elif value in {"dimensions", "dimension"}:
            mapping["dimension"] = col
        elif value == "videolength":
            mapping["video_length"] = col
        elif value in {"vastorvpaid", "vastvpaid"}:
            mapping["vast_vpaid"] = col
        elif value == "adname":
            mapping["ad"] = col
        elif value == "traffickingnotes":
            mapping["notes"] = col
        elif value == "action":
            mapping["action"] = col
        elif value == "creativefilename":
            mapping["creative"] = col
        elif value.startswith("studiocreative"):
            mapping["studio"] = col
        elif value.startswith("rotation"):
            mapping["rotation"] = col
        elif value == "startdate":
            mapping["start"] = col
        elif value == "enddate":
            mapping["end"] = col
        elif value.startswith("clickthroughurl"):
            mapping["url"] = col

    required = {
        "site", "placement", "dimension", "ad", "action", "creative",
        "rotation", "start", "end", "url",
    }
    missing = sorted(required - set(mapping))

    if missing:
        raise ValueError(
            "Traffic_Doc template is missing required columns: "
            + ", ".join(missing)
        )

    return mapping


def _snapshot_row_format(sheet, row_number: int) -> dict:
    cells = {}
    for col in range(1, sheet.max_column + 1):
        cell = sheet.cell(row_number, col)
        cells[col] = {
            "style": copy(cell._style),
            "number_format": cell.number_format,
        }

    return {
        "cells": cells,
        "height": sheet.row_dimensions[row_number].height,
    }


def _apply_row_format(sheet, row_number: int, snapshot: dict) -> None:
    for col, style in snapshot["cells"].items():
        cell = sheet.cell(row_number, col)
        cell._style = copy(style["style"])
        cell.number_format = style["number_format"]

    sheet.row_dimensions[row_number].height = snapshot["height"]


def _clear_values(sheet, start_row: int) -> None:
    for row in sheet.iter_rows(
        min_row=start_row,
        max_row=max(sheet.max_row, start_row + 5000),
        min_col=1,
        max_col=max(sheet.max_column, 30),
    ):
        for cell in row:
            cell.value = None


def _to_date(value):
    if value is None or value == "":
        return ""

    if isinstance(value, (datetime, date)):
        return value

    text = _clean(value)
    for fmt in ("%m/%d/%Y", "%m-%d-%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass

    return text


def generate_bfas_tsheet(
    placement_text: str,
    creative_files,
    url_mapping_text: str,
    start_date=None,
    end_date=None,
    campaign_name: str = "",
    site_name: str = "Nexxen",
    template_path: str | Path | None = None,
) -> tuple[bytes, list[str], dict]:
    """
    Generate BFAS Traffic_Doc.

    BFAS rules:
    - No Prisma upload.
    - Placement Name = Ad Name.
    - Creative match requires dimension/video-length match, then shared words.
    - One URL may apply to all creatives, or URLs can be mapped by set keyword.
    - Rotation = 100%.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("openpyxl is required to generate BFAS T-sheets.") from exc

    template = Path(template_path) if template_path else MASTER_TEMPLATE

    if not template.exists():
        raise FileNotFoundError(
            f"BFAS template not found: {template.name}"
        )

    if not _clean(placement_text):
        raise ValueError("Please provide at least one BFAS Placement Name.")

    preview = preview_bfas_setup(
        placement_text=placement_text,
        creative_files=creative_files,
        url_mapping_text=url_mapping_text,
    )

    if not preview["creative_names"]:
        raise ValueError("No supported BFAS creative files were found.")

    workbook = load_workbook(template, keep_vba=True)

    if TRAFFIC_SHEET not in workbook.sheetnames:
        raise KeyError(f"Missing worksheet: {TRAFFIC_SHEET}")

    sheet = workbook[TRAFFIC_SHEET]

    # Optional campaign name.
    if campaign_name:
        sheet["B1"] = campaign_name

    header_row = _find_header_row(sheet)
    first_data_row = header_row + 1
    columns = _header_map(sheet, header_row)

    style_snapshot = _snapshot_row_format(sheet, first_data_row)
    _clear_values(sheet, first_data_row)

    warnings = list(preview["warnings"])

    for index, item in enumerate(preview["rows"]):
        row = first_data_row + index
        _apply_row_format(sheet, row, style_snapshot)

        placement = item["placement_name"]
        size = item["size"]
        creative = item["creative"]
        url = item["url"]

        is_video = bool(_extract_video_length(placement))
        display_dimension = _extract_dimension(placement)

        values = {
            "site": _clean(site_name),
            "placement": placement,
            "dimension": display_dimension or size,
            "ad": placement,
            "action": "New",
            "creative": creative,
            "rotation": 1,
            "start": _to_date(start_date),
            "end": _to_date(end_date),
            "url": url,
        }

        if "studio" in columns:
            values["studio"] = "N"

        if "additional_pixels" in columns:
            values["additional_pixels"] = ""

        if "video_length" in columns:
            video_length = _extract_video_length(placement)
            values["video_length"] = (
                int(video_length[:-1]) if video_length else ""
            )

        if "vast_vpaid" in columns:
            values["vast_vpaid"] = "Vpaid" if is_video else ""

        for key, value in values.items():
            if key in columns:
                sheet.cell(row=row, column=columns[key]).value = value

        sheet.cell(row=row, column=columns["rotation"]).number_format = "0%"

        if "start" in columns:
            sheet.cell(row=row, column=columns["start"]).number_format = "mm/dd/yyyy"
        if "end" in columns:
            sheet.cell(row=row, column=columns["end"]).number_format = "mm/dd/yyyy"

    # BFAS is single-creative in this workflow; clear stale Multi rows.
    if MULTI_SHEET in workbook.sheetnames:
        multi = workbook[MULTI_SHEET]
        first_multi_data_row = 2
        for merged in list(multi.merged_cells.ranges):
            if merged.min_row >= first_multi_data_row:
                multi.unmerge_cells(str(merged))

        for row in multi.iter_rows(
            min_row=first_multi_data_row,
            max_row=max(multi.max_row, 5000),
            min_col=1,
            max_col=max(multi.max_column, 20),
        ):
            for cell in row:
                cell.value = None

    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except Exception:
        pass

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    return output.getvalue(), warnings, preview
