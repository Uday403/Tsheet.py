from __future__ import annotations

import csv
import io
import re
import zipfile
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
MASTER_TEMPLATE = BASE_DIR / "master_template.xlsm"

PRISMA_SHEET = "Prisma Export - Paste as values"
TRAFFIC_SHEET = "Traffic_Doc"
MULTI_SHEET = "Multi-Ad or Creative Rotation"
SEE_MULTI = "See Multi-Ad tab"
TRACKING_1X1 = "Tracking_1x1"


def _clean(v) -> str:
    return "" if v is None else str(v).strip()


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def _read_uploaded_bytes(uploaded_file) -> bytes:
    uploaded_file.seek(0)
    return uploaded_file.read()


def _read_csv_rows(uploaded_file) -> list[list[str]]:
    text = _decode(_read_uploaded_bytes(uploaded_file))
    try:
        dialect = csv.Sniffer().sniff(text[:10000], delimiters=",;\t|")
        reader = csv.reader(io.StringIO(text), dialect)
    except csv.Error:
        reader = csv.reader(io.StringIO(text))
    return [list(r) for r in reader]


def _campaign_name(rows: list[list[str]]) -> str:
    for row in rows[:20]:
        if row and _clean(row[0]).lower().startswith("campaign name"):
            return _clean(row[1] if len(row) > 1 else "")
    return ""


def _find_header_row(rows: list[list[str]]) -> int:
    for i, row in enumerate(rows):
        vals = {_clean(x).lower() for x in row}
        if "ad server id" in vals and "placement name" in vals and ("tag size" in vals or "ad size" in vals):
            return i
    raise ValueError("Could not find the Prisma placement header row.")


def _colmap(header: list[str]) -> dict[str, int]:
    return {_clean(v).lower(): i for i, v in enumerate(header)}


def _parse_date(v):
    if isinstance(v, datetime):
        return v
    s = _clean(v)
    if not s:
        return None
    for fmt in (
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y",
        "%d/%m/%Y", "%d-%m-%Y", "%m-%d-%Y",
    ):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return s


def _normalize_size(v: str) -> str:
    s = _clean(v).lower().replace(" ", "").replace("×", "x")
    if s in {"nanxnan", "0x0", "nan", ""}:
        return "1x1"
    m = re.search(r"(\d+)x(\d+)", s)
    return f"{m.group(1)}x{m.group(2)}" if m else s


def _extract_concept(campaign: str) -> str:
    # BRSP_DPS_F26 Cascadia H2T_Perf -> H2T
    m = re.search(r"\b([A-Za-z0-9]+)_Perf\b", campaign, flags=re.I)
    if m:
        return m.group(1)
    parts = [p for p in re.split(r"[_\s]+", campaign) if p]
    for i, p in enumerate(parts):
        if p.lower().startswith("perf") and i:
            return parts[i - 1]
    return "BROOKS"


def _publisher_label(placement_name: str, supplier_name: str) -> str:
    parts = [p.strip() for p in placement_name.split("_")]
    if len(parts) >= 4 and parts[3]:
        return parts[3].title()
    supplier = _clean(supplier_name)
    if "outside" in supplier.lower():
        return "Outside"
    words = supplier.split()
    return words[0].title() if words else "Publisher"


def _creative_names(creative_files) -> list[str]:
    names: list[str] = []
    for f in creative_files or []:
        name = getattr(f, "name", "") or ""
        data = _read_uploaded_bytes(f)
        if name.lower().endswith(".zip"):
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    names.extend(Path(n).name for n in z.namelist() if not n.endswith("/"))
            except zipfile.BadZipFile:
                names.append(name)
        else:
            names.append(name)

    result, seen = [], set()
    for n in names:
        stem = Path(n).stem.strip()
        if stem and stem.lower() not in seen:
            seen.add(stem.lower())
            result.append(stem)
    return result


def _matches_dimension(name: str, size: str) -> bool:
    return re.search(rf"(?<!\d){re.escape(size)}(?!\d)", name, flags=re.I) is not None


def _extract_urls(text: str) -> list[str]:
    urls = re.findall(r"https?://[^\s<>\"']+", text or "", flags=re.I)
    out, seen = [], set()
    for u in urls:
        u = u.rstrip(".,;)")
        if u not in seen:
            seen.add(u)
            out.append(u)
    return out


def _url_tokens(value: str) -> set[str]:
    value = re.sub(r"https?://[^/]+", " ", value.lower())
    value = value.split("?", 1)[0]
    tokens = set(re.findall(r"[a-z]+\d*", value))
    stop = {"www", "brooksrunning", "com", "en", "us", "featured", "displayads", "paid", "outside", "mag", "ooh", "f26"}
    return {t for t in tokens if t not in stop and len(t) > 2}


def _choose_url(creative: str, urls: list[str]) -> tuple[str, str | None]:
    """Return (url, warning). Uses concept/slug matching, never row order."""
    if not urls:
        return "", None
    if len(urls) == 1:
        return urls[0], None

    c = creative.lower()
    # Strong Brooks sample rule: Cascadia creative -> Cascadia destination.
    if "cascadia" in c:
        hits = [u for u in urls if "cascadia" in u.lower()]
        if len(hits) == 1:
            return hits[0], None
    else:
        non_cascadia = [u for u in urls if "cascadia" not in u.lower()]
        if len(non_cascadia) == 1:
            return non_cascadia[0], None

    ct = _url_tokens(creative)
    scored = []
    for u in urls:
        score = len(ct & _url_tokens(u))
        scored.append((score, u))
    scored.sort(reverse=True)
    if scored and scored[0][0] > 0 and (len(scored) == 1 or scored[0][0] > scored[1][0]):
        return scored[0][1], None
    return "", f"Could not uniquely map a URL to creative: {creative}"


def _find_sheet_header(ws, required: list[str], scan_rows: int = 20) -> tuple[int, dict[str, int]]:
    required_norm = [x.lower().strip() for x in required]
    for r in range(1, min(ws.max_row, scan_rows) + 1):
        mapping = {}
        for c in range(1, ws.max_column + 1):
            val = _clean(ws.cell(r, c).value).lower()
            if val:
                mapping[val] = c
        if all(any(req == key or req in key for key in mapping) for req in required_norm):
            result = {}
            for req in required_norm:
                for key, col in mapping.items():
                    if req == key or req in key:
                        result[req] = col
                        break
            return r, result
    raise ValueError(f"Could not find required headers in sheet {ws.title}: {required}")


def _find_col(ws, header_row: int, names: list[str]) -> int | None:
    for c in range(1, ws.max_column + 1):
        val = _clean(ws.cell(header_row, c).value).lower()
        for name in names:
            n = name.lower()
            if val == n or n in val:
                return c
    return None


def _copy_row_style(ws, source_row: int, target_row: int, max_col: int | None = None):
    max_col = max_col or ws.max_column
    for c in range(1, max_col + 1):
        src, dst = ws.cell(source_row, c), ws.cell(target_row, c)
        if src.has_style:
            dst._style = copy(src._style)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
        dst.number_format = src.number_format
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def _write_prisma(ws, rows: list[list[str]]):
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(r_idx, c_idx).value = value


def _clear_values(ws, start_row: int, end_row: int):
    for r in range(start_row, end_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).value = None


def preview_brooks_setup(prisma_file, creative_files, url_mapping_text: str = ""):
    rows = _read_csv_rows(prisma_file)
    campaign = _campaign_name(rows)
    concept = _extract_concept(campaign)
    h = _find_header_row(rows)
    cols = _colmap(rows[h])
    creatives = _creative_names(creative_files)
    urls = _extract_urls(url_mapping_text)

    required = ["ad server id", "placement name", "media outlet / supplier name (prisma)", "flight start date", "flight end date"]
    missing = [x for x in required if x not in cols]
    if "tag size" not in cols and "ad size" not in cols:
        missing.append("tag size/ad size")
    if missing:
        raise ValueError("Missing Prisma columns: " + ", ".join(missing))

    placements, warnings = [], []
    direct = multi = one_by_one = unmatched = 0

    for row in rows[h + 1:]:
        if not any(_clean(x) for x in row):
            continue

        def get(name):
            idx = cols.get(name)
            return row[idx] if idx is not None and idx < len(row) else ""

        pid, pname = _clean(get("ad server id")), _clean(get("placement name"))
        if not pid and not pname:
            continue

        size = _normalize_size(get("tag size") or get("ad size"))
        publisher = _publisher_label(pname, get("media outlet / supplier name (prisma)"))
        ad_name = f"{concept} - {publisher} - {size}"
        matches = [] if size == "1x1" else [n for n in creatives if _matches_dimension(n, size)]

        if size == "1x1":
            status, one_by_one = "Tracking 1x1", one_by_one + 1
        elif len(matches) == 0:
            status, unmatched = "No creative match", unmatched + 1
            warnings.append(f"{pid or pname}: no creative matched {size}.")
        elif len(matches) == 1:
            status, direct = "Direct", direct + 1
        else:
            status, multi = f"Multi ({len(matches)})", multi + 1

        placements.append({
            "placement_id": pid, "placement_name": pname, "size": size,
            "publisher": publisher, "ad_name": ad_name,
            "start_date": _parse_date(get("flight start date")),
            "end_date": _parse_date(get("flight end date")),
            "matches": matches, "status": status,
        })

    return {
        "campaign": campaign, "concept": concept, "placements": placements,
        "creative_count": len(creatives), "url_count": len(urls),
        "direct_count": direct, "multi_count": multi,
        "tracking_1x1_count": one_by_one, "unmatched_count": unmatched,
        "warnings": warnings,
    }


def generate_brooks_tsheet(
    prisma_file,
    creative_files,
    url_mapping_text: str = "",
    apply_dynata_display: bool = False,
    template_path: str | Path | None = None,
):
    """Generate Brooks T-sheet while preserving the master template layout/styles."""
    preview = preview_brooks_setup(prisma_file, creative_files, url_mapping_text)
    rows = _read_csv_rows(prisma_file)
    urls = _extract_urls(url_mapping_text)

    template = Path(template_path) if template_path else MASTER_TEMPLATE
    if not template.exists():
        raise FileNotFoundError(f"Template not found: {template}")

    wb = load_workbook(template, keep_vba=True)
    for sheet in (PRISMA_SHEET, TRAFFIC_SHEET, MULTI_SHEET):
        if sheet not in wb.sheetnames:
            raise ValueError(f"Template is missing required sheet: {sheet}")

    prisma_ws, traffic_ws, multi_ws = wb[PRISMA_SHEET], wb[TRAFFIC_SHEET], wb[MULTI_SHEET]
    _write_prisma(prisma_ws, rows)

    # Keep template formulas/labels. Only set default base URL when the cell exists.
    if traffic_ws["B5"].value is not None or _clean(traffic_ws["A5"].value).lower().startswith("default click"):
        traffic_ws["B5"] = "https://www.brooksrunning.com"

    # Detect headers dynamically. This fixes the previous two-blank-row problem.
    traffic_header, _ = _find_sheet_header(traffic_ws, ["ad name", "creative file name", "start date", "end date"])
    multi_header, _ = _find_sheet_header(multi_ws, ["ad name", "creative file name", "start date", "end date"])
    traffic_start = traffic_header + 1
    multi_start = multi_header + 1

    # Header-driven columns prevent the Multi dates/creative fields from shifting.
    tc = {
        "pixels": _find_col(traffic_ws, traffic_header, ["additional pixels"]),
        "site": _find_col(traffic_ws, traffic_header, ["site name"]),
        "pid": _find_col(traffic_ws, traffic_header, ["dcm placement id"]),
        "placement": _find_col(traffic_ws, traffic_header, ["placement name"]),
        "size": _find_col(traffic_ws, traffic_header, ["dimensions"]),
        "ad": _find_col(traffic_ws, traffic_header, ["ad name"]),
        "notes": _find_col(traffic_ws, traffic_header, ["trafficking notes"]),
        "action": _find_col(traffic_ws, traffic_header, ["action"]),
        "creative": _find_col(traffic_ws, traffic_header, ["creative file name"]),
        "studio": _find_col(traffic_ws, traffic_header, ["studio creative"]),
        "rotation": _find_col(traffic_ws, traffic_header, ["rotation %"]),
        "start": _find_col(traffic_ws, traffic_header, ["start date"]),
        "end": _find_col(traffic_ws, traffic_header, ["end date"]),
        "url": _find_col(traffic_ws, traffic_header, ["click through url", "click-through url"]),
    }
    mc = {
        "ad": _find_col(multi_ws, multi_header, ["ad name"]),
        "action": _find_col(multi_ws, multi_header, ["action"]),
        "notes": _find_col(multi_ws, multi_header, ["trafficking notes"]),
        "creative": _find_col(multi_ws, multi_header, ["creative file name"]),
        "studio": _find_col(multi_ws, multi_header, ["studio creative"]),
        "rotation": _find_col(multi_ws, multi_header, ["rotation %"]),
        "start": _find_col(multi_ws, multi_header, ["start date"]),
        "end": _find_col(multi_ws, multi_header, ["end date"]),
        "url": _find_col(multi_ws, multi_header, ["click through url", "click-through url"]),
    }
    missing_tc = [k for k, v in tc.items() if k not in {"pixels", "notes"} and v is None]
    missing_mc = [k for k, v in mc.items() if k != "notes" and v is None]
    if missing_tc or missing_mc:
        raise ValueError(f"Template header mapping failed. Traffic: {missing_tc}; Multi: {missing_mc}")

    # Capture a formatted template data row BEFORE clearing it, then reuse it for every output row.
    traffic_style_row = traffic_start
    multi_style_row = multi_start
    traffic_style = [copy(traffic_ws.cell(traffic_style_row, c)._style) for c in range(1, traffic_ws.max_column + 1)]
    multi_style = [copy(multi_ws.cell(multi_style_row, c)._style) for c in range(1, multi_ws.max_column + 1)]
    traffic_height = traffic_ws.row_dimensions[traffic_style_row].height
    multi_height = multi_ws.row_dimensions[multi_style_row].height

    _clear_values(traffic_ws, traffic_start, max(traffic_ws.max_row, traffic_start + len(preview["placements"]) + 20))
    _clear_values(multi_ws, multi_start, max(multi_ws.max_row, multi_start + 300))

    warnings = list(preview["warnings"])
    multi_row = multi_start

    def apply_style(ws, row, styles, height):
        for c, style in enumerate(styles, start=1):
            ws.cell(row, c)._style = copy(style)
        ws.row_dimensions[row].height = height

    for offset, p in enumerate(preview["placements"]):
        r = traffic_start + offset
        apply_style(traffic_ws, r, traffic_style, traffic_height)
        is_1x1, matches = p["size"] == "1x1", p["matches"]

        def setv(col, value):
            if col:
                traffic_ws.cell(r, col).value = value

        setv(tc["pixels"], "N" if is_1x1 else ("Y" if apply_dynata_display else "N"))
        setv(tc["site"], "OUTSIDE INTEGRATED MEDIA" if p["publisher"] == "Outside" else p["publisher"])
        setv(tc["pid"], int(p["placement_id"]) if p["placement_id"].isdigit() else p["placement_id"])
        setv(tc["placement"], p["placement_name"])
        setv(tc["size"], p["size"])
        setv(tc["ad"], p["ad_name"])
        setv(tc["notes"], "Please add 2026 Dynata Pixel" if (apply_dynata_display and not is_1x1) else None)
        setv(tc["action"], "New")
        setv(tc["studio"], "N")
        setv(tc["start"], p["start_date"])
        setv(tc["end"], p["end_date"])

        if is_1x1:
            setv(tc["creative"], TRACKING_1X1)
            setv(tc["rotation"], "100%")
            # If only one URL is supplied, it is safe to use it for 1x1. Otherwise leave for review.
            if len(urls) == 1:
                setv(tc["url"], urls[0])
        elif len(matches) == 1:
            creative = matches[0]
            setv(tc["creative"], creative)
            setv(tc["rotation"], "100%")
            url, warning = _choose_url(creative, urls)
            setv(tc["url"], url or None)
            if warning:
                warnings.append(warning)
        elif len(matches) >= 2:
            setv(tc["creative"], None)
            setv(tc["rotation"], SEE_MULTI)
            setv(tc["url"], SEE_MULTI)
            for creative in matches:
                apply_style(multi_ws, multi_row, multi_style, multi_height)
                vals = {
                    "ad": p["ad_name"], "action": "New", "notes": None,
                    "creative": creative, "studio": "N", "rotation": "Even",
                    "start": p["start_date"], "end": p["end_date"],
                }
                url, warning = _choose_url(creative, urls)
                vals["url"] = url or None
                if warning:
                    warnings.append(warning)
                for key, value in vals.items():
                    if mc.get(key):
                        multi_ws.cell(multi_row, mc[key]).value = value
                multi_row += 1

        for col in (tc["start"], tc["end"]):
            if col:
                traffic_ws.cell(r, col).number_format = "mm/dd/yyyy"

    for r in range(multi_start, multi_row):
        for col in (mc["start"], mc["end"]):
            if col:
                multi_ws.cell(r, col).number_format = "mm/dd/yyyy"

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue(), warnings, preview
